from django.shortcuts import render, redirect, get_object_or_404
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from .models import LessonPlan, LessonStepTemplate, Subject, LessonType
from django.core.files.base import ContentFile
from django.db.models import QuerySet
from .forms import (LessonPlanFileUploadForm, SubjectModelForm, LessonPlanModelForm,
                    SubjectTypeModelForm, LessonStepTemplateModelForm, EditLessonStepTemplateModelForm,
                    LessonPlanFilterForm)
from django.views.generic.edit import FormMixin
from django.urls import reverse_lazy, reverse
from django.http.response import HttpResponseBadRequest
import os.path
from django.http import HttpRequest
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http.response import HttpResponseRedirect

from docx.section import Section
from docx import Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.enum.section import WD_ORIENT
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO


def generate_plan_lesson_word(lesson_plan):
    """
    Генерация документа word плана-конспекта урока с таблицей и информацией об уроке
    """
    steps_lesson = get_lesson_plan_steps(lesson_plan)
    document: Document = Document()
    section = get_section(document, 0)
    set_section_horizontal_orient(section)

    heading = gat_heading_lvl_0(document, 'План-конспект урока', 0)
    set_paragraph_alignment(heading, WD_ALIGN_PARAGRAPH.CENTER)

    set_document_font_name(document, "Times New Roman")
    set_paragraph_line_spacing(document.paragraphs[0], 1.5)

    lesson_fields = {"Предмет": lesson_plan.subject.name,
                     "Тема урока": lesson_plan.topic,
                     "Тип урока": lesson_plan.subject_type.name,
                     "Класс": lesson_plan.grade.name,
                     "Цель урока": lesson_plan.goal,
                     "Оборудование": lesson_plan.equipment.name}

    for i, field in enumerate(lesson_fields, 1):
        if lesson_fields[field]:
            document.add_paragraph(f"{field}: {lesson_fields[field]}")
            document.paragraphs[i].runs[0].font.size = Pt(14)
            document.paragraphs[i].runs[0].font.bold = True

    table_rows = len(steps_lesson) + 1
    table = get_table_document(document, rows_count=table_rows, columns_count=3)
    table.style = 'Table Grid'

    table.columns[0].width = section.page_width // 3 // 2
    table.columns[1].width = int((section.page_width - table.columns[0].width) / 2.5)
    table.columns[2].width = int((section.page_width - table.columns[0].width) / 2.5)

    hdr_cells = get_head_table(table)
    head_table_cells_text = ('Этап урока', 'Деятельность учителя', 'Деятельность обучающихся')
    set_style_head_table_lesson_plan(hdr_cells, head_table_cells_text)
    set_text_column_steps_lesson(table, steps_lesson)

    buffer = BytesIO()
    save_document_in_buffer(document, buffer)
    filename = f"lesson_plan_{lesson_plan.pk}.docx"
    django_file = ContentFile(buffer.read(), filename)
    lesson_plan.plan_file.save(filename, django_file)


def get_lesson_plan_steps(lesson_plan: LessonPlan) -> QuerySet:
    return LessonStepTemplate.object_step.filter(lesson_type=lesson_plan.subject_type).order_by("order")


def get_section(document: Document, num: int) -> Section:
    return document.sections[num]


def set_section_horizontal_orient(section: Section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)  # ~A4 в альбомной (297 мм)
    section.page_height = Inches(8.27)  # ~A4 в портретной (210 мм)


def gat_heading_lvl_0(document: Document, text, level) -> Paragraph:
    return document.add_heading(text, level)


def set_paragraph_alignment(paragraph: Paragraph, alignment) -> None:
    paragraph.paragraph_format.alignment = alignment


def set_document_font_name(document: Document, font_name: str) -> None:
    document.paragraphs[0].runs[0].font.name = font_name


def set_paragraph_line_spacing(paragraph: Paragraph, line_spacing: float | int) -> None:
    paragraph.paragraph_format.line_spacing = line_spacing


def get_table_document(document, rows_count: int, columns_count: int) -> Table:
    return document.add_table(rows=rows_count, cols=columns_count)


def get_head_table(table: Table) -> tuple[_Cell]:
    return table.rows[0].cells


def set_style_head_table_lesson_plan(head: tuple[_Cell], head_cells_text: tuple[str, ...]) -> None:
    for i, text in enumerate(head_cells_text):
        head[i].text = text
        head_paragraph: Paragraph = head[i].paragraphs[0]
        set_paragraph_alignment(head_paragraph, WD_ALIGN_PARAGRAPH.CENTER)
        set_paragraph_line_spacing(head_paragraph, 1)
        head_paragraph.runs[0].font.bold = True
        head_paragraph.runs[0].font.size = Pt(12)


def set_text_column_steps_lesson(table: Table, steps_lesson: QuerySet) -> None:
    for i, step in enumerate(steps_lesson, 1):
        table.rows[i].cells[0].text = f"{i}. {step.name}"


def save_document_in_buffer(document: Document, buffer: BytesIO) -> None:
    document.save(buffer)
    buffer.seek(0)


def docx_to_html(path: str) -> str | None:
    """
    Мз docx файла парсит текст таблицы и оборачивает в теги html
    """
    document = Document(path)
    table_html = "<table class='table table-bordered m-0'>"
    if document.tables:
        table = document.tables[0]
        for i, row in enumerate(table.rows, 1):
            if i == 1:
                table_html += "<thead class='text-center'>"
            if i == 2:
                table_html += "<tbody>"
            table_html += "<tr>"
            for cell in row.cells:
                if i == 1:
                    table_html += "".join(f"<th>{cell.text}</th>")
                else:
                    table_html += "".join(f"<td>{cell.text}</td>")
            table_html += "</tr>"
            if i == 1:
                table_html += "</thead>"
        table_html += "</tbody></table>"
        return table_html
    return None


class LessonPlanListView(ListView):
    model = LessonPlan
    template_name = "lesson_planner/index.html"
    context_object_name = "lesson_plans"
    paginate_by = 6

    def get_context_data(self, *, object_list=None, **kwargs):
        context = super().get_context_data(object_list=None, **kwargs)
        context["title"] = "Главная страница"
        context["form"] = LessonPlanFilterForm(self.request.GET)
        return context

    def get_queryset(self) -> QuerySet:
        queryset = super().get_queryset()
        if self.request.user.is_authenticated:
            queryset = queryset.filter(creator=self.request.user)
        form = LessonPlanFilterForm(self.request.GET)
        subjects = grades = types_lessons = None
        if form.is_valid():
            subjects = form.cleaned_data.get("subjects")
            grades = form.cleaned_data.get("grades")
            types_lessons = form.cleaned_data.get("types_lessons")
            date_from = form.cleaned_data.get("date_from")
            date_before = form.cleaned_data.get("date_before")
        if subjects:
            queryset = queryset.filter(subject__in=subjects)
        if grades:
            queryset = queryset.filter(grade__in=grades)
        if types_lessons:
            queryset = queryset.filter(subject_type__in=types_lessons)
        if date_from:
            queryset = queryset.filter(date_created__gte=date_from)
        if date_before:
            queryset = queryset.filter(date_created__lte=date_before)
        return queryset


class LessonPlanDetailView(FormMixin, LoginRequiredMixin, DetailView):
    model = LessonPlan
    context_object_name = "lesson_plan"
    template_name = "lesson_planner/lessonplan_detail.html"
    form_class = LessonPlanFileUploadForm

    def __init__(self, **kwargs):
        super().__init__()
        self.object = None

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        print(self.object.plan_file)
        context['title'] = "План урока"
        context['form'] = LessonPlanFileUploadForm(instance=self.object)
        if self.object.plan_file and os.path.exists(self.object.plan_file.path):
            context["docx_html"] = docx_to_html(self.object.plan_file.path)
        return context

    def post(self, request: HttpRequest, *args, **kwargs):
        self.object: LessonPlan = self.get_object()
        action = request.POST.get("action")

        if action == "generate":
            generate_plan_lesson_word(lesson_plan=self.object)
            return redirect(self.object.get_absolute_url())
        elif action == "upload":
            form = LessonPlanFileUploadForm(request.POST, request.FILES, instance=self.object)
            if form.is_valid():
                form.save()
                return redirect(self.object.get_absolute_url())
        elif action == "delete":
            if self.object.plan_file:
                self.object.plan_file.delete(save=False)
                self.object.plan_file = None
                self.object.save()
                return redirect(self.object.get_absolute_url())
        return self.get(request, args, kwargs)


class LessonPlanCreateView(LoginRequiredMixin, CreateView):
    model = LessonPlan
    success_url = reverse_lazy("lesson_planner:index")
    form_class = LessonPlanModelForm

    def form_valid(self, form: LessonPlanModelForm) -> HttpResponseRedirect:
        form.instance.creator = self.request.user
        return super().form_valid(form)


class LessonPlanDeleteView(LoginRequiredMixin, DeleteView):
    model = LessonPlan
    success_url = reverse_lazy("lesson_planner:index")

    def __init__(self, **kwargs):
        super().__init__()
        self.object = None

    def post(self, request, *args, **kwargs):
        self.object: LessonPlan = self.get_object()
        self.object.delete()
        return redirect(self.success_url)


class LessonPlanUpdateView(LoginRequiredMixin, UpdateView):
    model = LessonPlan
    form_class = LessonPlanModelForm

    def get_success_url(self):
        return reverse("lesson_planner:lesson_detail", args=[self.object.pk])


class SubjectListView(FormMixin, LoginRequiredMixin, ListView):
    model = Subject
    context_object_name = "subjects"
    form_class = SubjectModelForm

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['title'] = 'Предметы'
        context['form'] = self.get_form(form_class=self.form_class)
        return context

    def post(self, request: HttpRequest, *args, **kwargs):
        subject_id = request.POST.get("subject_id")
        if subject_id != "add":
            try:
                subject: Subject = self.model.objects.get(id=subject_id)
            except self.model.DoesNotExist:
                return HttpResponseBadRequest("Предмет не найден")
            form = SubjectModelForm(request.POST, instance=subject)
        else:
            form = SubjectModelForm(request.POST)
        if form.is_valid():
            subject = form.save(commit=False)
            subject.creator = request.user
            subject.save()
            return redirect(request.path)
        return self.get(request, args, kwargs)

    def get_queryset(self) -> QuerySet:
        queryset = super().get_queryset()
        queryset = queryset.filter(creator=self.request.user)
        return queryset


class SubjectDeleteView(LoginRequiredMixin, DeleteView):
    model = Subject
    success_url = reverse_lazy("lesson_planner:subjects_list")


class SubjectTypeListView(FormMixin, LoginRequiredMixin, ListView):
    model = LessonType
    context_object_name = "subject_types"
    form_class = SubjectTypeModelForm

    def get_context_data(self, *, object_list=None, **kwargs):
        context = super().get_context_data(object_list=None, **kwargs)
        context["title"] = "Типы уроков"
        context["form_add_type"] = self.get_form(form_class=self.form_class)
        # steps_lessons = self.model.objects.prefetch_related("step_templates")
        # context["steps_lessons"] = {}
        # for lesson_type in steps_lessons:
        #     for step in lesson_type.step_templates.all():
        #         if lesson_type == step.lesson_type:
        #             context["steps_lessons"][lesson_type] = context["steps_lessons"].setdefault(lesson_type, []) + [step]
        return context

    def post(self, request: HttpRequest, *args, **kwargs):
        subject_type_id = request.POST.get("subject_type_id")
        if subject_type_id == "add_type":
            form = SubjectTypeModelForm(request.POST)
            if form.is_valid():
                subject_type: LessonType = form.save(commit=False)
                subject_type.creator = request.user
                subject_type.save()
                return redirect(request.path)
        return self.get(request, args, kwargs)

    def get_queryset(self):
        queryset = super().get_queryset()
        return queryset.filter(creator=self.request.user)


class SubjectTypeDeleteView(LoginRequiredMixin, DeleteView):
    model = LessonType
    success_url = reverse_lazy("lesson_planner:subject_types_list")


class LessonStepTemplateListView(FormMixin, LoginRequiredMixin, ListView):
    model = LessonStepTemplate
    context_object_name = "lesson_steps"
    form_class = LessonStepTemplateModelForm

    def get_queryset(self) -> QuerySet:
        # queryset = super().get_queryset()
        # filters = {
        #     LessonStepTemplate.lesson_type_id: self.kwargs["pk"],
        #     LessonStepTemplate.lesson_type: self.request.user
        # }
        # queryset = queryset.filter(filters)
        return LessonStepTemplate.object_step.filter(lesson_type_id=self.kwargs["pk"]) # lesson_type_id=self.kwargs["pk"]

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["title"] = "Этапы урока"
        context["lesson_type"] = get_object_or_404(LessonType, pk=self.kwargs["pk"])
        context["form_add_step"] = self.get_form(form_class=self.form_class)
        context["form_edit_step"] = EditLessonStepTemplateModelForm()
        return context

    def post(self, request: HttpRequest, *args, **kwargs):
        subject_step_id = request.POST.get("subject_step_id")
        if subject_step_id == "add_step":
            form = LessonStepTemplateModelForm(request.POST)
            if form.is_valid():
                step = form.save(commit=False)
                lesson_type = get_object_or_404(LessonType, pk=self.kwargs["pk"])
                step.lesson_type = lesson_type
                step.save()
                return redirect(request.path)
        else:
            step: LessonStepTemplate = self.get_queryset().get(id=subject_step_id)
            form = EditLessonStepTemplateModelForm(request.POST, instance=step)
            if form.is_valid():
                form.save()
                return redirect(request.path)
            print(form.errors)
        return self.get(request, args, kwargs)


class LessonStepTemplateDeleteView(LoginRequiredMixin, DeleteView):
    model = LessonStepTemplate

    def get_success_url(self):
        return reverse("lesson_planner:lessons_steps_list", args=[self.object.lesson_type.pk])
